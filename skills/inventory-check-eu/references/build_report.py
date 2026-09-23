"""Builds the EU Inventory Check .xlsx report from a captured BigQuery row set.

Battle-tested, exact-structure build script for the inventory-check-eu skill. Do not
rewrite or reshape this logic inline in a one-off script -- if something about the
output needs to change, fix it here and re-run, the same way presentation-creator's
slide-templates.js is treated.

Usage inside the claude.ai sandbox:
    1. Write the captured row data (list of dicts, one per result row, keys matching
       HEADER below) to /home/claude/rows.json
    2. Copy this file to /home/claude/build_report.py
    3. Run: python /home/claude/build_report.py
    4. The output lands at /mnt/user-data/outputs/eu_inventory_check_<today>.xlsx
    5. The script prints a summary dict -- use those counts for the chat summary
       (see Core Capability 3 in SKILL.md for the exact wording)
"""

import json
import statistics
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROWS_PATH = Path("/home/claude/rows.json")
OUTPUT_DIR = Path("/mnt/user-data/outputs")

# Exact header row -- order is locked, matches the team's reference template
# (eu_inventory_check.xlsx) column-for-column. Never reorder without an explicit
# requester override.
HEADER = [
    "asin", "sku", "sku_type", "eu_inventory", "gb_inventory",
    "GB", "DE", "ES", "FR", "IT", "NL", "BE", "IE", "PL", "SE", "TR",
]

# Country columns start at F (index 6, 1-based) and end at P (index 16).
COUNTRY_COL_START = 6
COUNTRY_COL_END = 16
ALL_COUNTRIES = HEADER[COUNTRY_COL_START - 1:COUNTRY_COL_END]
PAN_EU_COUNTRIES = ["DE", "ES", "FR", "IT", "NL"]
# All 7 of these share EUR -- safe to compare raw prices directly. GB (GBP),
# PL (PLN), SE (SEK), TR (TRY) are excluded from price-outlier comparison
# specifically, since comparing un-converted currencies produces meaningless
# "outliers" (confirmed: they accounted for 1,543 of 1,590 flags in testing,
# purely from currency, not real pricing errors). Deliberately no hardcoded
# exchange rates here -- those go stale; see the skill standard's anti-pattern
# on time-sensitive facts stated as permanent truth.
EUROZONE_COUNTRIES = ["DE", "ES", "FR", "IT", "NL", "BE", "IE"]
PRICE_OUTLIER_THRESHOLD = 0.5  # 50% deviation from a SKU's median price

COUNTRY_COLUMN_WIDTH = 18
ROW_HEIGHT = 21
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)
HEADER_FONT = Font(bold=True)

# Vertical middle for the whole table. Horizontal: eu_inventory/gb_inventory
# centered, country columns (F:P) right-aligned, everything else left/default.
ALIGN_DEFAULT = Alignment(vertical="center")
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")
CENTERED_COLUMNS = {"eu_inventory", "gb_inventory"}

# Fixed widths for the non-country columns (A-E). Country columns (F-P) all
# use COUNTRY_COLUMN_WIDTH instead.
COLUMN_WIDTHS = {
    "asin": 15,
    "sku": 40,
    "sku_type": 15,
    "eu_inventory": 15,
    "gb_inventory": 15,
}

# Analysis sheet section styling -- same banner/header convention already
# used by the pivot-fields-from-powerbi skill's report sheet.
SECTION_BANNER_FILL = PatternFill(fgColor="FF1F3864", fill_type="solid")
SECTION_BANNER_FONT = Font(color="FFFFFFFF", bold=True, size=12)
COLUMN_HEADER_FILL = PatternFill(fgColor="FFD9D9D9", fill_type="solid")
COLUMN_HEADER_FONT = Font(bold=True)


def build_main_sheet(wb: Workbook, rows: list[dict]) -> None:
    ws = wb.active
    ws.title = "main"

    ws.append(HEADER)
    for row in rows:
        ws.append([row.get(col) for col in HEADER])

    last_row = len(rows) + 1  # +1 for the header row
    country_range = f"F2:P{last_row}"

    for name, width in COLUMN_WIDTHS.items():
        col = HEADER.index(name) + 1
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = width

    for col in range(COUNTRY_COL_START, COUNTRY_COL_END + 1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = COUNTRY_COLUMN_WIDTH

    for r in range(1, last_row + 1):
        ws.row_dimensions[r].height = ROW_HEIGHT
        for c in range(1, len(HEADER) + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = THIN_BORDER
            if c >= COUNTRY_COL_START:
                cell.alignment = ALIGN_RIGHT
            elif HEADER[c - 1] in CENTERED_COLUMNS:
                cell.alignment = ALIGN_CENTER
            else:
                cell.alignment = ALIGN_DEFAULT

    for c in range(1, len(HEADER) + 1):
        ws.cell(row=1, column=c).font = HEADER_FONT

    sku_col = HEADER.index("sku") + 1
    for r in range(1, last_row + 1):
        ws.cell(row=r, column=sku_col).font = HEADER_FONT

    # Applied last so it overrides the per-column alignment above: every
    # header cell is centered both horizontally and vertically, regardless
    # of what its column's data rows use.
    for c in range(1, len(HEADER) + 1):
        ws.cell(row=1, column=c).alignment = ALIGN_CENTER

    # Freeze row 1 so it stays visible when scrolling, and add filter
    # dropdown arrows on the header row (equivalent to Ctrl+Shift+L).
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:P{last_row}"

    # Rule 1: literal 0 -> red font on red fill
    ws.conditional_formatting.add(
        country_range,
        CellIsRule(
            operator="equal",
            formula=["0"],
            font=Font(color="FF9C0006"),
            fill=PatternFill(bgColor="FFC00000"),
        ),
    )

    # Rule 2: contains "Inactive" -> yellow fill
    ws.conditional_formatting.add(
        country_range,
        FormulaRule(
            formula=['NOT(ISERROR(SEARCH("Inactive",F2)))'],
            fill=PatternFill(bgColor="FFFFFF00"),
        ),
    )

    # Rule 3: contains "Active" -> green fill
    ws.conditional_formatting.add(
        country_range,
        FormulaRule(
            formula=['NOT(ISERROR(SEARCH("Active",F2)))'],
            fill=PatternFill(bgColor="FF00B050"),
        ),
    )


# ======================================================
# ANALYSIS -- data-health findings
# ======================================================

def parse_cell(value):
    """Splits a 'Status | Price' country cell into (status, price).

    price is a float, or None if missing/unparseable. Returns (None, None)
    for an empty cell (no listing in that country at all).
    """
    if not value:
        return None, None
    parts = value.split(" | ", 1)
    status = parts[0] if parts else None
    price = None
    if len(parts) > 1:
        try:
            price = float(parts[1])
        except ValueError:
            price = None
    return status, price


def find_pan_eu_gaps(rows: list[dict]):
    """EU-prefixed SKUs missing a price in any of DE/ES/FR/IT/NL.

    Split HIGH (has EU pooled inventory -- stock exists but can't be sold in
    that country) vs LOW (no EU inventory either way).
    """
    high, low = [], []
    for row in rows:
        sku = row.get("sku") or ""
        if not sku.startswith("EU-"):
            continue
        missing = [c for c in PAN_EU_COUNTRIES if not row.get(c)]
        if not missing:
            continue
        eu_inv = float(row.get("eu_inventory") or 0)
        record = {
            "sku": sku, "asin": row.get("asin"),
            "missing_countries": ", ".join(missing),
            "eu_inventory": eu_inv,
        }
        (high if eu_inv > 0 else low).append(record)
    return high, low


def find_gb_gaps(rows: list[dict]):
    """GB-prefixed SKUs with no GB price, split by gb_inventory."""
    high, low = [], []
    for row in rows:
        sku = row.get("sku") or ""
        if not sku.startswith("GB-") or row.get("GB"):
            continue
        gb_inv = float(row.get("gb_inventory") or 0)
        record = {"sku": sku, "asin": row.get("asin"), "gb_inventory": gb_inv}
        (high if gb_inv > 0 else low).append(record)
    return high, low


def find_missing_by_country(rows: list[dict]):
    """Per country (all 11, all SKUs): every SKU with no price/listing at all
    in that country. Returns (summary dict country->count, detail list).
    """
    summary = {c: 0 for c in ALL_COUNTRIES}
    detail = []
    for row in rows:
        for c in ALL_COUNTRIES:
            if not row.get(c):
                summary[c] += 1
                detail.append({"country": c, "sku": row.get("sku"), "asin": row.get("asin")})
    return summary, detail


def find_stock_zero_listings(rows: list[dict]):
    """SKUs with EU or GB inventory but no listing in any of the 11 countries
    -- stock with no way to sell it anywhere. Likely the most serious finding.
    """
    out = []
    for row in rows:
        eu_inv = float(row.get("eu_inventory") or 0)
        gb_inv = float(row.get("gb_inventory") or 0)
        if eu_inv <= 0 and gb_inv <= 0:
            continue
        if any(row.get(c) for c in ALL_COUNTRIES):
            continue
        out.append({
            "sku": row.get("sku"), "asin": row.get("asin"),
            "eu_inventory": eu_inv, "gb_inventory": gb_inv,
        })
    return out


def find_all_inactive(rows: list[dict]):
    """SKUs with at least one listing, but every listing is Inactive (none
    Active). Split HIGH (has EU or GB inventory) / LOW (none).
    """
    high, low = [], []
    for row in rows:
        statuses, listed = [], []
        for c in ALL_COUNTRIES:
            status, _ = parse_cell(row.get(c))
            if status:
                statuses.append(status)
                listed.append(c)
        if not statuses or any(s == "Active" for s in statuses):
            continue
        eu_inv = float(row.get("eu_inventory") or 0)
        gb_inv = float(row.get("gb_inventory") or 0)
        record = {
            "sku": row.get("sku"), "asin": row.get("asin"),
            "listed_countries": ", ".join(listed),
            "eu_inventory": eu_inv, "gb_inventory": gb_inv,
        }
        (high if (eu_inv > 0 or gb_inv > 0) else low).append(record)
    return high, low


def find_price_outliers(rows: list[dict], threshold: float = PRICE_OUTLIER_THRESHOLD):
    """Same SKU priced very differently across Eurozone countries -- flags any
    country whose price deviates more than `threshold` (50% default) from the
    median of that SKU's own listed prices. Possible pricing data-entry error.

    Scoped to EUROZONE_COUNTRIES only (all share EUR) -- comparing raw prices
    across different currencies (GB/PL/SE/TR) produces meaningless results.
    """
    out = []
    for row in rows:
        prices = {}
        for c in EUROZONE_COUNTRIES:
            _, price = parse_cell(row.get(c))
            if price is not None:
                prices[c] = price
        if len(prices) < 2:
            continue
        median = statistics.median(prices.values())
        if median == 0:
            continue
        for c, price in prices.items():
            deviation = abs(price - median) / median
            if deviation > threshold:
                out.append({
                    "sku": row.get("sku"), "asin": row.get("asin"), "country": c,
                    "price": price, "median_price": round(median, 2),
                    "deviation_pct": round(deviation * 100, 1),
                })
    return out


def find_active_zero_inventory(rows: list[dict]):
    """SKU shows Active somewhere but eu_inventory AND gb_inventory are both
    0 -- possible stockout/overselling risk (could be normal if sold through
    a channel outside this pooled inventory -- flagged for review, not
    necessarily an error).
    """
    out = []
    for row in rows:
        eu_inv = float(row.get("eu_inventory") or 0)
        gb_inv = float(row.get("gb_inventory") or 0)
        if eu_inv > 0 or gb_inv > 0:
            continue
        active = [c for c in ALL_COUNTRIES if parse_cell(row.get(c))[0] == "Active"]
        if active:
            out.append({"sku": row.get("sku"), "asin": row.get("asin"), "active_countries": ", ".join(active)})
    return out


def _write_section(ws, start_row: int, title: str, columns: list[str], records: list[dict]) -> int:
    """Writes one banner + column-header + data-rows section starting at
    start_row. Returns the row number the next section should start at.
    """
    ncols = len(columns)
    ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=max(ncols, 2))
    banner = ws.cell(row=start_row, column=1, value=f"{title} ({len(records)})")
    banner.fill = SECTION_BANNER_FILL
    banner.font = SECTION_BANNER_FONT

    header_row = start_row + 1
    for i, col_name in enumerate(columns, start=1):
        cell = ws.cell(row=header_row, column=i, value=col_name)
        cell.fill = COLUMN_HEADER_FILL
        cell.font = COLUMN_HEADER_FONT

    r = header_row + 1
    if not records:
        ws.cell(row=r, column=1, value="(none found)")
        r += 1
    else:
        for record in records:
            for i, col_name in enumerate(columns, start=1):
                ws.cell(row=r, column=i, value=record.get(col_name))
            r += 1

    return r + 2  # 2 blank rows before the next section


def build_analysis_sheet(wb: Workbook, rows: list[dict]) -> dict:
    ws = wb.create_sheet("Analysis")
    r = 1

    pan_eu_high, pan_eu_low = find_pan_eu_gaps(rows)
    r = _write_section(ws, r, "Pan-EU price gaps -- HAS EU inventory (HIGH priority)",
                        ["sku", "asin", "missing_countries", "eu_inventory"], pan_eu_high)
    r = _write_section(ws, r, "Pan-EU price gaps -- no EU inventory (LOW priority)",
                        ["sku", "asin", "missing_countries", "eu_inventory"], pan_eu_low)

    gb_high, gb_low = find_gb_gaps(rows)
    r = _write_section(ws, r, "GB price gap -- HAS GB inventory (HIGH priority)",
                        ["sku", "asin", "gb_inventory"], gb_high)
    r = _write_section(ws, r, "GB price gap -- no GB inventory (LOW priority)",
                        ["sku", "asin", "gb_inventory"], gb_low)

    by_country_summary, by_country_detail = find_missing_by_country(rows)
    summary_records = [{"country": c, "missing_sku_count": n} for c, n in by_country_summary.items()]
    r = _write_section(ws, r, "Missing price by country -- summary",
                        ["country", "missing_sku_count"], summary_records)
    r = _write_section(ws, r, "Missing price by country -- detail",
                        ["country", "sku", "asin"], by_country_detail)

    stock_zero_listings = find_stock_zero_listings(rows)
    r = _write_section(ws, r, "Stock with zero listings anywhere",
                        ["sku", "asin", "eu_inventory", "gb_inventory"], stock_zero_listings)

    inactive_high, inactive_low = find_all_inactive(rows)
    r = _write_section(ws, r, "All listings inactive -- HAS inventory (HIGH priority)",
                        ["sku", "asin", "listed_countries", "eu_inventory", "gb_inventory"], inactive_high)
    r = _write_section(ws, r, "All listings inactive -- no inventory (LOW priority)",
                        ["sku", "asin", "listed_countries", "eu_inventory", "gb_inventory"], inactive_low)

    price_outliers = find_price_outliers(rows)
    r = _write_section(ws, r, "Price outliers, Eurozone only (>50% deviation from SKU's median price)",
                        ["sku", "asin", "country", "price", "median_price", "deviation_pct"], price_outliers)

    active_zero_inventory = find_active_zero_inventory(rows)
    r = _write_section(ws, r, "Active listing but zero pooled inventory",
                        ["sku", "asin", "active_countries"], active_zero_inventory)

    ws.column_dimensions["A"].width = 40
    for letter in "BCDEF":
        ws.column_dimensions[letter].width = 20

    return {
        "pan_eu_gaps_high": len(pan_eu_high),
        "pan_eu_gaps_low": len(pan_eu_low),
        "gb_gaps_high": len(gb_high),
        "gb_gaps_low": len(gb_low),
        "missing_by_country": by_country_summary,
        "stock_zero_listings": len(stock_zero_listings),
        "all_inactive_high": len(inactive_high),
        "all_inactive_low": len(inactive_low),
        "price_outliers": len(price_outliers),
        "active_zero_inventory": len(active_zero_inventory),
    }


def build(rows: list[dict]) -> tuple[Path, dict]:
    wb = Workbook()
    build_main_sheet(wb, rows)
    summary = build_analysis_sheet(wb, rows)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    output_path = OUTPUT_DIR / f"eu_inventory_check_{date.today().isoformat()}.xlsx"
    wb.save(output_path)
    return output_path, summary


if __name__ == "__main__":
    rows = json.loads(ROWS_PATH.read_text(encoding="utf-8"))
    path, summary = build(rows)
    print(f"Wrote {len(rows)} rows to {path}")
    print(json.dumps(summary, indent=2))
